"""
Document Ingestion & Text Processing Service.
Handles extraction from PDF, DOCX, and TXT files, content cleaning,
hash generation, chunking with configurable overlap, and page-level metadata.
"""

import hashlib
import os
import re
import unicodedata
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
try:
    import pypdf as pdf_module
except ImportError:
    import PyPDF2 as pdf_module
import docx

from backend.app.core.logging import logger


class DocumentProcessingService:
    """Enterprise document extractor and text normalization pipeline."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
    MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB limit

    @staticmethod
    def compute_sha256(file_path: Path) -> str:
        """Computes SHA-256 hash of a file for exact duplicate detection."""
        hasher = hashlib.sha256()
        with open(file_path, "rb") as f:
            while chunk := f.read(65536):
                hasher.update(chunk)
        return hasher.hexdigest()

    @staticmethod
    def clean_text(text: str) -> str:
        """Removes unreadable control characters, standardizes whitespace while preserving paragraphs and Unicode NFKC fidelity."""
        if not text:
            return ""
        # Unicode normalization (NFKC) for uniform glyphs, ligatures, and compatibility forms
        text = unicodedata.normalize('NFKC', text)
        # Remove zero-width spaces, joiners, and non-printable control chars except newlines and tabs
        text = re.sub(r'[\u200b-\u200d\ufeff]', '', text)
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        # Normalize carriage returns
        text = text.replace('\r\n', '\n').replace('\r', '\n')
        # Compress multiple blank lines to double newline
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    def extract_text(self, file_path: Path, extension: str) -> List[Dict[str, Any]]:
        """
        Extracts structured text page-by-page or section-by-section.
        Returns a list of dicts: [{"page_number": int, "text": str}]
        """
        ext = extension.lower()
        if ext == ".pdf":
            return self._extract_pdf(file_path)
        elif ext == ".docx":
            return self._extract_docx(file_path)
        elif ext == ".txt":
            return self._extract_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {extension}. Supported formats: PDF, DOCX, TXT")

    def _extract_pdf(self, file_path: Path) -> List[Dict[str, Any]]:
        pages = []
        try:
            with open(file_path, "rb") as f:
                reader = pdf_module.PdfReader(f)
                num_pages = len(reader.pages)
                if num_pages == 0:
                    raise ValueError("PDF document is empty (0 pages).")
                for page_idx, page in enumerate(reader.pages):
                    raw_text = page.extract_text() or ""
                    cleaned = self.clean_text(raw_text)
                    pages.append({
                        "page_number": page_idx + 1,
                        "text": cleaned
                    })
                if not any(p["text"] for p in pages):
                    raise ValueError("PDF document contains no extractable text. Scanned image-only PDFs require OCR processing.")
        except Exception as e:
            logger.error(f"Failed to extract PDF text from {file_path.name}: {e}")
            raise ValueError(f"Corrupted or unreadable PDF: {str(e)}")
        return pages

    def _extract_docx(self, file_path: Path) -> List[Dict[str, Any]]:
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract from tables as well
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n\n".join(paragraphs)
            cleaned = self.clean_text(full_text)
            if not cleaned:
                raise ValueError("DOCX document contains no extractable text.")
            
            # Estimate pages assuming ~400 words per page
            words = cleaned.split()
            page_size_words = 400
            total_pages = max(1, (len(words) + page_size_words - 1) // page_size_words)
            
            pages = []
            for p_idx in range(total_pages):
                start = p_idx * page_size_words
                end = start + page_size_words
                page_content = " ".join(words[start:end])
                pages.append({
                    "page_number": p_idx + 1,
                    "text": page_content
                })
            return pages
        except Exception as e:
            logger.error(f"Failed to extract DOCX text from {file_path.name}: {e}")
            raise ValueError(f"Corrupted or unreadable DOCX: {str(e)}")

    def _extract_txt(self, file_path: Path) -> List[Dict[str, Any]]:
        try:
            # Try utf-8 first, fallback to latin-1
            content = ""
            for encoding in ("utf-8", "latin-1", "cp1252"):
                try:
                    with open(file_path, "r", encoding=encoding) as f:
                        content = f.read()
                    break
                except UnicodeDecodeError:
                    continue

            cleaned = self.clean_text(content)
            if not cleaned:
                raise ValueError("TXT file is empty.")

            # Chunk into simulated pages of 400 words
            words = cleaned.split()
            page_size_words = 400
            total_pages = max(1, (len(words) + page_size_words - 1) // page_size_words)
            pages = []
            for p_idx in range(total_pages):
                start = p_idx * page_size_words
                end = start + page_size_words
                pages.append({
                    "page_number": p_idx + 1,
                    "text": " ".join(words[start:end])
                })
            return pages
        except Exception as e:
            logger.error(f"Failed to extract TXT text from {file_path.name}: {e}")
            raise ValueError(f"Unreadable TXT file: {str(e)}")

    def chunk_document(
        self,
        pages: List[Dict[str, Any]],
        chunk_size: int = 500,
        chunk_overlap: int = 80
    ) -> List[Dict[str, Any]]:
        """
        Splits extracted pages into overlapping chunks while preserving page provenance.
        `chunk_size` and `chunk_overlap` measured in words for robust semantic cohesion.
        """
        chunks = []
        global_chunk_idx = 0

        for page in pages:
            page_num = page["page_number"]
            page_text = page["text"]
            if not page_text:
                continue

            words = page_text.split()
            if not words:
                continue

            if len(words) <= chunk_size:
                chunks.append({
                    "chunk_index": global_chunk_idx,
                    "page_number": page_num,
                    "content": page_text,
                    "token_count": len(words)
                })
                global_chunk_idx += 1
            else:
                step = chunk_size - chunk_overlap
                if step <= 0:
                    step = chunk_size // 2

                for i in range(0, len(words), step):
                    chunk_words = words[i:i + chunk_size]
                    if not chunk_words:
                        continue
                    chunk_text = " ".join(chunk_words)
                    chunks.append({
                        "chunk_index": global_chunk_idx,
                        "page_number": page_num,
                        "content": chunk_text,
                        "token_count": len(chunk_words)
                    })
                    global_chunk_idx += 1
                    if i + chunk_size >= len(words):
                        break

        return chunks


document_processor = DocumentProcessingService()
